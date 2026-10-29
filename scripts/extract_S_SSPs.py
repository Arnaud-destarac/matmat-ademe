"""
Extrait, pour chaque scénario IMAGE de data/SSP/Scenarios, les facteurs F, x et S
2050/2020 par région, et les écrit dans Matching_extensions_EXIOBASE_IMAGE.xlsx,
une feuille par scénario (feuille effacée puis recréée si elle existe déjà).

Pour chaque ligne de la feuille "Matching" :
- F 2050/2020 : ratio 2050/2020 de la formule "Numérateur F", évaluée séparément
  en 2020 et en 2050. La formule combine des variables IMAGE avec les opérateurs
  "+", "-" (ou "−"), "x", "/", des nombres et des parenthèses ; les parenthèses
  placées après une variable sont des unités et sont ignorées, ex.
  "iNH3_volatization_arable (kgN/an) x (1 − Land Cover|Cropland|Energy Crops / Land Cover|Cropland)".
  "max (A, B)" / "min (A, B)" s'appliquent année par année (ex. plancher à 0).
  "A OU B" (priorité la plus faible) vaut A, sauf si A est négatif en 2020 ou en
  2050 pour la région, auquel cas B est utilisé (ex. P - agriculture - soil/water).
  Chaque variable est lue dans Nutrients.xlsx (scénario "PB_<scénario>") pour le
  domaine Nutrients, dans le csv du scénario sinon ; si elle est absente de cette
  source, elle est lue dans l'autre (ex. iNH3_volatization_* du domaine
  Air_emissions, Land Cover|* dans les lignes du domaine Nutrients).
- x 2050/2020 : ratio de la formule "Dénominateur x" (csv du scénario).
- S 2050/2020 : F/x ; si le numérateur est de la forme "1 / <formule>" (rendements),
  F est le ratio de 1 / <formule>, soit S ; 1 si F ou x n'est pas calculable.
  S vaut aussi 1 quand l'intensité 2020 de la région (F/x) est négligeable devant celle
  du Monde (moins de MIN_RELATIVE_INTENSITY fois) : le ratio 2050/2020 d'une valeur
  2020 quasi nulle n'a pas de sens (ex. charbon dans le tertiaire indonésien,
  F 2020 ~ 1e-26 x Monde -> S ~ 1e15), et ces facteurs, appliqués aux coefficients de
  A par generate_shock_dS.py, rendent la matrice de Leontief inutilisable.
- Si un rendement "Yield|..." manque dans le scénario (ex. Yield|Oilcrops dans SSP1),
  la ligne reprend les variables F et x de la ligne "Cropland".

Baseline commune : les valeurs 2020 de toutes les variables (csv et Nutrients.xlsx)
sont celles du scénario BASELINE_SCENARIO (SSP2) ; seules les valeurs 2050 sont
propres à chaque scénario. Une variable absente de SSP2 garde sa valeur 2020.

Les régions IMAGE sont agrégées selon la colonne "Code agrégé" de la feuille
"Régions" (somme des valeurs 2020 et 2050 avant calcul des ratios). Les valeurs World
absentes ou nulles sont remplacées par la somme des régions.

La S Monde de chaque scénario est aussi écrite dans la feuille Matching, colonne
"Facteur S 2050/2020 (Monde <scénario>)" (régénérée à chaque exécution), ainsi que
F et x Monde du scénario de référence, colonnes "F 2050/2020 (Monde SSP2)" et
"x 2050/2020 (Monde SSP2)".
"""

import re
import sys
from copy import copy
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import column_index_from_string, get_column_letter

MODULE_PLANEFR_DIR = Path(r"C:\Users\Arnaud\Documents\CIRED\PlaneFR\Code\Module_PlaneFR")

sys.path.insert(0, str(MODULE_PLANEFR_DIR / "notebooks"))
from planefr_lib import config  # noqa: E402

SSP_DIR = config.DATA_DIR / "SSP"
SCENARIOS_DIR = SSP_DIR / "Scenarios"
NUTRIENTS_FILE = SSP_DIR / "Nutrients.xlsx"
MATCHING_FILE = SSP_DIR / "Matching_extensions_EXIOBASE_IMAGE.xlsx"

MATCHING_SHEET = "Matching"
REGIONS_SHEET = "Régions"
NUTRIENTS_DOMAIN = "Nutrients"
YIELD_FALLBACK_EXTENSION = "Cropland"

LABEL_COLUMNS = ["Domaine", "Extension EXIOBASE", "Secteurs EXIOBASE"]

YEAR_BASE, YEAR_TARGET = 2020, 2050
BASELINE_SCENARIO = "SSP2"  # fournit les valeurs 2020 (F et x) de tous les scénarios
VALUE_COLUMNS = ["F 2050/2020", "x 2050/2020", "S 2050/2020"]
# Intensité 2020 (F/x) de la région, rapportée à celle du Monde, en dessous de laquelle S vaut 1
MIN_RELATIVE_INTENSITY = 1e-2

# Colonnes "Facteur S 2050/2020 (Monde <scénario>)" de la feuille Matching
WORLD_REGION, WORLD_LABEL = "World", "Monde"
WORLD_S_HEADER = "Facteur S 2050/2020"
# Colonnes F et x Monde du scénario de référence de la feuille Matching (indice dans les résultats)
WORLD_BASELINE_COLUMNS = {f"F 2050/2020 ({WORLD_LABEL} {BASELINE_SCENARIO})": 0,
                          f"x 2050/2020 ({WORLD_LABEL} {BASELINE_SCENARIO})": 1}

# Opérateurs des formules (entourés d'espaces dans les cellules), par priorité croissante.
OPERATOR_SYMBOLS = {"OU": "OU", "+": "+", "-": "-", "−": "-", "x": "*", "×": "*", "*": "*", "/": "/"}
OPERATOR_PATTERN = "|".join(re.escape(op) for op in OPERATOR_SYMBOLS)
PRECEDENCE = [("OU",), ("+", "-"), ("*", "/")]
ARITHMETIC = {"+": np.add, "-": np.subtract, "*": np.multiply, "/": np.divide}
# Fonctions à deux arguments séparés par une virgule (avec ou sans espace ; une virgule entre
# deux chiffres est une virgule décimale), appliquées année par année, ex. "max (A, 0)".
FUNCTIONS = {"max": np.maximum, "min": np.minimum}
EMPTY_CELLS = ("", "–", "-")


def strip_units(text):
    """Retire les parenthèses d'unité placées après une variable, ex. "A (kgN/an) x (1 − B)" -> "A x (1 − B)".

    Une parenthèse est un regroupement si elle contient un opérateur ou suit un opérateur,
    une autre parenthèse ouvrante, un nom de fonction ou le début de la cellule.
    """
    def replace(match):
        prefix = text[:match.start()].rstrip()
        is_group = (re.search(rf"\s({OPERATOR_PATTERN})\s", match.group(1))
                    or not prefix or prefix.endswith("(")
                    or re.search(rf"(^|\s)({OPERATOR_PATTERN})$", prefix)
                    or re.search(rf"(^|[\s(])({'|'.join(FUNCTIONS)})$", prefix))
        return match.group(0) if is_group else ""
    return re.sub(r"\s*\(([^()]*)\)", replace, text)


def tokenize(text):
    """Jetons ("name", texte), ("op", symbole) ou ("paren", "(" / ")" / ",")."""
    tokens = []
    for i, part in enumerate(re.split(rf"\s+({OPERATOR_PATTERN})\s+|([()]|,(?!\d)|(?<!\d),)", text)):
        if part is None or not part.strip():
            continue
        kind = ("name", "op", "paren")[i % 3]
        tokens.append((kind, OPERATOR_SYMBOLS[part] if kind == "op" else part.strip()))
    return tokens


def parse_expression(tokens, pos=0, level=0):
    """Analyse descendante : renvoie (arbre, position suivante).

    Arbre : ("var", nom), ("num", valeur) ou (opérateur ou fonction, gauche, droite).
    """
    if level == len(PRECEDENCE):
        return parse_atom(tokens, pos)
    left, pos = parse_expression(tokens, pos, level + 1)
    while pos < len(tokens) and tokens[pos][0] == "op" and tokens[pos][1] in PRECEDENCE[level]:
        right, next_pos = parse_expression(tokens, pos + 1, level + 1)
        left, pos = (tokens[pos][1], left, right), next_pos
    return left, pos


def parse_atom(tokens, pos):
    if pos >= len(tokens):
        raise ValueError("formule incomplète")
    kind, value = tokens[pos]
    if (kind, value) == ("paren", "("):
        node, pos = parse_expression(tokens, pos + 1)
        if pos >= len(tokens) or tokens[pos] != ("paren", ")"):
            raise ValueError("parenthèse non fermée")
        return node, pos + 1
    if kind == "name" and value in FUNCTIONS and tokens[pos + 1:pos + 2] == [("paren", "(")]:
        first, pos = parse_expression(tokens, pos + 2)
        if tokens[pos:pos + 1] != [("paren", ",")]:
            raise ValueError(f"{value} attend deux arguments")
        second, pos = parse_expression(tokens, pos + 1)
        if tokens[pos:pos + 1] != [("paren", ")")]:
            raise ValueError("parenthèse non fermée")
        return (value, first, second), pos + 1
    if kind == "name":
        try:
            return ("num", float(value.replace(",", "."))), pos + 1  # décimale française "0,1"
        except ValueError:
            return ("var", value), pos + 1
    raise ValueError(f"jeton inattendu '{value}'")


def parse_formula(cell):
    """Analyse une cellule "Numérateur F" / "Dénominateur x".

    Renvoie (arbre, inverse) ; arbre None si pas de variable.
    Ex. "A + B (Mt MS/an)" -> (("+", A, B), False) ; "1 / Yield|Cereal (t/ha)" -> (("var", "Yield|Cereal"), True).
    """
    if pd.isna(cell):
        return None, False
    text = strip_units(str(cell).strip()).strip()
    if text in EMPTY_CELLS:
        return None, False
    tokens = tokenize(text)
    try:
        tree, pos = parse_expression(tokens)
        if pos != len(tokens):
            raise ValueError(f"jeton inattendu '{tokens[pos][1]}'")
    except ValueError as error:
        raise ValueError(f"Formule invalide ({error}) : {cell!r}") from None
    if tree[0] == "/" and tree[1] == ("num", 1.0):
        return tree[2], True
    return tree, False


def formula_variables(tree):
    """Variables IMAGE de la formule, dans l'ordre d'apparition."""
    if tree is None or tree[0] == "num":
        return []
    if tree[0] == "var":
        return [tree[1]]
    return list(dict.fromkeys(formula_variables(tree[1]) + formula_variables(tree[2])))


def load_matching():
    df = pd.read_excel(MATCHING_FILE, sheet_name=MATCHING_SHEET)
    return df.rename(columns={
        df.columns[0]: "Domaine",
        df.columns[1]: "Extension EXIOBASE",
        df.columns[2]: "Secteurs EXIOBASE",
        df.columns[3]: "Numérateur F",
        df.columns[4]: "Dénominateur x",
    })


def load_region_maps():
    """Dictionnaires code IMAGE (csv) -> code agrégé et code IMAGE-GNM -> code agrégé."""
    df = pd.read_excel(MATCHING_FILE, sheet_name=REGIONS_SHEET).dropna(subset=["Code agrégé"])
    gnm_col, image_col, agg_col = df.columns[0], df.columns[1], "Code agrégé"
    image_map = dict(zip(df[image_col], df[agg_col]))
    gnm_map = dict(zip(df[gnm_col], df[agg_col]))
    region_order = list(dict.fromkeys(df[agg_col]))
    return image_map, gnm_map, region_order


def aggregate(df, region_map):
    """Tableau (variable x région agrégée) des valeurs 2020 et 2050, sommées par région agrégée."""
    df = df.assign(Region=df["Region"].map(region_map)).dropna(subset=["Region"])
    return with_world_fallback(df.groupby(["Variable", "Region"])[[YEAR_BASE, YEAR_TARGET]].sum(min_count=1))


def with_world_fallback(data):
    """Valeurs World absentes ou nulles remplacées par la somme des autres régions.

    Le csv ne renseigne pas toujours World (ex. Production|Pulp and Paper|Volume = 0,
    Secondary Energy|Heat absent) alors que les régions le sont.
    """
    is_world = data.index.get_level_values("Region") == WORLD_REGION
    world = data[is_world].droplevel("Region")
    region_sum = data[~is_world].groupby(level="Variable").sum(min_count=1)
    world = world.reindex(world.index.union(region_sum.index))
    world = world.mask(world == 0).combine_first(region_sum).combine_first(world)
    world.index = pd.MultiIndex.from_product([world.index, [WORLD_REGION]], names=["Variable", "Region"])
    return pd.concat([data[~is_world], world]).sort_index()


def load_scenario(csv_path, image_map):
    df = pd.read_csv(csv_path, encoding="latin-1")
    df = df.rename(columns={str(YEAR_BASE): YEAR_BASE, str(YEAR_TARGET): YEAR_TARGET})
    return aggregate(df, image_map)


def load_nutrients(scenario_name, gnm_map):
    df = pd.read_excel(NUTRIENTS_FILE, header=1)
    df = df[df["Scenario"] == f"PB_{scenario_name}"]
    if df.empty:
        print(f"[info] Scénario PB_{scenario_name} absent de Nutrients.xlsx.")
    return aggregate(df, gnm_map)


def with_common_baseline(data, baseline):
    """Copie de data dont les valeurs 2020 sont celles du scénario de référence (si disponibles)."""
    data = data.copy()
    data[YEAR_BASE] = baseline[YEAR_BASE].reindex(data.index).combine_first(data[YEAR_BASE])
    return data


def has_variables(data, variables):
    known = data.index.get_level_values("Variable")
    return all(v in known for v in variables)


def f_sources(row, scenario_data, nutrients_data):
    """Sources des variables de F, par ordre de préférence : Nutrients.xlsx d'abord pour le domaine Nutrients."""
    if row["Domaine"] == NUTRIENTS_DOMAIN:
        return [nutrients_data, scenario_data]
    return [scenario_data, nutrients_data]


def lookup_values(sources, variable, region):
    """Valeurs [2020, 2050] de la variable dans la première source qui la contient (NaN sinon)."""
    for data in sources:
        try:
            return data.loc[(variable, region), [YEAR_BASE, YEAR_TARGET]].to_numpy(dtype=float)
        except KeyError:
            continue
    return np.full(2, np.nan)


def evaluate(tree, sources, region):
    """Valeurs [2020, 2050] de la formule pour une région (NaN si non calculable)."""
    kind = tree[0]
    if kind == "num":
        return np.full(2, tree[1])
    if kind == "var":
        return lookup_values(sources, tree[1], region)
    left = evaluate(tree[1], sources, region)
    if kind == "OU":
        # A négatif (2020 ou 2050) : indicateur de repli B
        return evaluate(tree[2], sources, region) if (left < 0).any() else left
    with np.errstate(divide="ignore", invalid="ignore"):
        result = {**ARITHMETIC, **FUNCTIONS}[kind](left, evaluate(tree[2], sources, region))
    return np.where(np.isfinite(result), result, np.nan)


def ratio(tree, sources, region):
    """Ratio 2050/2020 de la formule pour une région (NaN si non calculable)."""
    if tree is None:
        return np.nan
    base, target = evaluate(tree, sources, region)
    if pd.isna(base) or pd.isna(target) or base == 0:
        return np.nan
    return target / base


def relative_intensity(f_tree, f_src, x_tree, x_src, region):
    """Intensité 2020 F/x de la région rapportée à celle du Monde (NaN si non calculable)."""
    if f_tree is None or x_tree is None or region == WORLD_REGION:
        return np.nan
    f_region, f_world = evaluate(f_tree, f_src, region)[0], evaluate(f_tree, f_src, WORLD_REGION)[0]
    x_region, x_world = evaluate(x_tree, x_src, region)[0], evaluate(x_tree, x_src, WORLD_REGION)[0]
    with np.errstate(divide="ignore", invalid="ignore"):
        rel = abs((f_region / f_world) / (x_region / x_world))
    return rel if np.isfinite(rel) else np.nan


def compute_row(row, scenario_data, nutrients_data, region):
    """(F, x, S, négligeable) ; négligeable = S forcé à 1 car l'intensité 2020 de la région est
    négligeable (voir MIN_RELATIVE_INTENSITY)."""
    f_tree, inverse = parse_formula(row["Numérateur F"])
    x_tree, _ = parse_formula(row["Dénominateur x"])
    f_src = f_sources(row, scenario_data, nutrients_data)
    f = ratio(f_tree, f_src, region)
    x = ratio(x_tree, [scenario_data], region)
    if inverse:
        # F = ratio de "1 / <formule>" (= S), cohérent avec le numérateur affiché
        s = 1 / f if pd.notna(f) and f != 0 else 1.0
        return (s if pd.notna(f) and f != 0 else f), x, s, False
    if pd.isna(f) or pd.isna(x) or x == 0:
        return f, x, 1.0, False
    if relative_intensity(f_tree, f_src, x_tree, [scenario_data], region) < MIN_RELATIVE_INTENSITY:
        return f, x, 1.0, True
    return f, x, f / x, False


def apply_yield_fallback(matching, scenario_data, scenario_name):
    """Lignes "1 / Yield|..." dont le rendement manque : reprend F et x de la ligne Cropland."""
    cropland = matching.loc[matching["Extension EXIOBASE"] == YIELD_FALLBACK_EXTENSION].iloc[0]
    matching = matching.copy()
    for i, row in matching.iterrows():
        f_tree, inverse = parse_formula(row["Numérateur F"])
        f_vars = formula_variables(f_tree)
        if inverse and not has_variables(scenario_data, f_vars):
            matching.loc[i, ["Numérateur F", "Dénominateur x"]] = cropland[["Numérateur F", "Dénominateur x"]].values
            print(f"[info] {scenario_name} : {', '.join(f_vars)} absent, "
                  f"'{row['Extension EXIOBASE']}' reprend la ligne '{YIELD_FALLBACK_EXTENSION}'.")
    return matching


def check_variables(matching, scenario_data, nutrients_data, scenario_name):
    """Signale les variables de la feuille Matching introuvables dans les données."""
    for _, row in matching.iterrows():
        f_tree, _ = parse_formula(row["Numérateur F"])
        x_tree, _ = parse_formula(row["Dénominateur x"])
        for tree, sources in [(f_tree, f_sources(row, scenario_data, nutrients_data)), (x_tree, [scenario_data])]:
            for v in formula_variables(tree):
                if not any(has_variables(source, [v]) for source in sources):
                    print(f"[info] {scenario_name} : variable '{v}' introuvable ({row['Extension EXIOBASE']}).")


def write_sheet(wb, sheet_name, matching, results, region_order):
    """(Re)crée la feuille : 2 lignes d'en-tête (code région / F, x, S), puis une ligne par extension."""
    if sheet_name in wb.sheetnames:
        index = wb.sheetnames.index(sheet_name)
        wb.remove(wb[sheet_name])
        ws = wb.create_sheet(sheet_name, index)
    else:
        ws = wb.create_sheet(sheet_name)

    bold, center = Font(bold=True), Alignment(horizontal="center", vertical="center")
    n_labels = len(LABEL_COLUMNS)
    for col, label in enumerate(LABEL_COLUMNS, start=1):
        ws.cell(row=1, column=col, value=label).font = bold
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    for i, region in enumerate(region_order):
        first_col = n_labels + 1 + i * len(VALUE_COLUMNS)
        ws.cell(row=1, column=first_col, value=region)
        ws.merge_cells(start_row=1, start_column=first_col, end_row=1, end_column=first_col + len(VALUE_COLUMNS) - 1)
        for j, label in enumerate(VALUE_COLUMNS):
            ws.cell(row=2, column=first_col + j, value=label)
    for cells in ws.iter_rows(min_row=1, max_row=2):
        for cell in cells:
            cell.font, cell.alignment = bold, center

    for r, (_, row) in enumerate(matching.iterrows(), start=3):
        for col, label in enumerate(LABEL_COLUMNS, start=1):
            ws.cell(row=r, column=col, value=row[label])
        for i, region in enumerate(region_order):
            for j, value in enumerate(results[(r - 3, region)]):
                ws.cell(row=r, column=n_labels + 1 + i * len(VALUE_COLUMNS) + j,
                        value=None if pd.isna(value) else float(value))

    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60
    ws.freeze_panes = "D3"


def shift_columns(ws, col, delta):
    """Insère (delta = 1) ou supprime (delta = -1) la colonne col, en décalant aussi les largeurs de colonnes
    (non gérées par openpyxl)."""
    widths = {column_index_from_string(letter): dim.width for letter, dim in ws.column_dimensions.items()}
    ws.insert_cols(col) if delta > 0 else ws.delete_cols(col)
    for index in sorted(widths, reverse=delta > 0):
        if index > col or (delta > 0 and index == col):
            ws.column_dimensions[get_column_letter(index + delta)].width = widths[index]


def add_world_columns(wb, excel_rows, world_s):
    """Écrit dans la feuille Matching une colonne "Facteur S 2050/2020 (Monde <scénario>)" par scénario,
    valeurs arrondies à 2 décimales. La colonne du scénario de référence est écrasée sur place (formules
    comprises) ; les autres sont supprimées puis réinsérées, dans l'ordre, juste après elle.

    excel_rows : ligne Excel de chaque extension ; world_s : {scénario: [S Monde par extension]}.
    """
    ws = wb[MATCHING_SHEET]
    is_world_s = lambda h: str(h).startswith(WORLD_S_HEADER)
    header_of = lambda name: f"{WORLD_S_HEADER} ({WORLD_LABEL} {name})"
    # Mise en forme reprise de la première colonne "Facteur S" existante
    template = next((cell for cell in ws[1] if is_world_s(cell.value)), None)
    template_format = None if template is None else (
        copy(template.font), copy(template.fill), copy(template.border), copy(template.alignment),
        ws.column_dimensions[template.column_letter].width)

    for name in world_s:
        headers = [cell.value for cell in ws[1]]
        if name != BASELINE_SCENARIO and header_of(name) in headers:
            shift_columns(ws, headers.index(header_of(name)) + 1, -1)

    for name, values in world_s.items():
        headers = [cell.value for cell in ws[1]]
        header = header_of(name)
        if header in headers:
            col, action = headers.index(header) + 1, "régénérée"
        else:
            col = max((i for i, h in enumerate(headers, start=1) if is_world_s(h)), default=ws.max_column) + 1
            shift_columns(ws, col, 1)
            cell = ws.cell(row=1, column=col, value=header)
            if template_format is not None:
                cell.font, cell.fill, cell.border, cell.alignment, width = template_format
                ws.column_dimensions[cell.column_letter].width = width
            action = "ajoutée"
        for r, value in zip(excel_rows, values):
            ws.cell(row=r, column=col, value=None if pd.isna(value) else round(float(value), 2))
        print(f"Colonne '{header}' {action} dans la feuille '{MATCHING_SHEET}'.")


def write_world_baseline_columns(wb, excel_rows, world_results):
    """Écrit F et x Monde du scénario de référence dans les colonnes existantes de la feuille Matching,
    valeurs arrondies à 2 décimales.

    world_results : [(F, x, S) Monde par extension] du scénario de référence.
    """
    ws = wb[MATCHING_SHEET]
    headers = [cell.value for cell in ws[1]]
    for header, k in WORLD_BASELINE_COLUMNS.items():
        if header not in headers:
            print(f"[info] Colonne '{header}' absente de la feuille '{MATCHING_SHEET}', non écrite.")
            continue
        col = headers.index(header) + 1
        for r, values in zip(excel_rows, world_results):
            ws.cell(row=r, column=col, value=None if pd.isna(values[k]) else round(float(values[k]), 2))
        print(f"Colonne '{header}' régénérée dans la feuille '{MATCHING_SHEET}'.")


def main():
    matching = load_matching().dropna(subset=["Extension EXIOBASE"])
    excel_rows = list(matching.index + 2)  # en-tête en ligne 1
    matching = matching.reset_index(drop=True)
    image_map, gnm_map, region_order = load_region_maps()
    wb = openpyxl.load_workbook(MATCHING_FILE)

    baseline_data = load_scenario(SCENARIOS_DIR / f"{BASELINE_SCENARIO}.csv", image_map)
    baseline_nutrients = load_nutrients(BASELINE_SCENARIO, gnm_map)

    world_s = {}
    for csv_path in sorted(SCENARIOS_DIR.glob("*.csv")):
        scenario_name = csv_path.stem
        scenario_data = with_common_baseline(load_scenario(csv_path, image_map), baseline_data)
        nutrients_data = with_common_baseline(load_nutrients(scenario_name, gnm_map), baseline_nutrients)
        scenario_matching = apply_yield_fallback(matching, scenario_data, scenario_name)
        check_variables(scenario_matching, scenario_data, nutrients_data, scenario_name)
        results = {
            (i, region): compute_row(row, scenario_data, nutrients_data, region)
            for i, row in scenario_matching.iterrows()
            for region in region_order
        }
        neglected = [(excel_rows[i], region) for (i, region), values in results.items() if values[3]]
        if neglected:
            print(f"[info] {scenario_name} : S forcé à 1 pour {len(neglected)} cellule(s) (intensité 2020 "
                  f"< {MIN_RELATIVE_INTENSITY:g} x Monde) : "
                  + ", ".join(f"ligne {r} {region}" for r, region in neglected))
        results = {key: values[:3] for key, values in results.items()}
        write_sheet(wb, scenario_name[:31], matching, results, region_order)
        print(f"Feuille '{scenario_name}' écrite.")
        world_s[scenario_name] = [results[(i, WORLD_REGION)][2] for i in range(len(matching))]
        if scenario_name == BASELINE_SCENARIO:
            write_world_baseline_columns(wb, excel_rows, [results[(i, WORLD_REGION)] for i in range(len(matching))])

    add_world_columns(wb, excel_rows, world_s)

    try:
        wb.save(MATCHING_FILE)
    except PermissionError:
        sys.exit(f"Impossible d'écrire {MATCHING_FILE} : fermer le fichier dans Excel puis relancer.")
    print(f"Classeur écrit : {MATCHING_FILE}")


if __name__ == "__main__":
    main()
